from docxtpl import DocxTemplate,InlineImage
from docx import Document
import io
doc = Document('XXFPGA测试记录.docx')
tables = doc.tables
cel = tables[0].cell(0,0)
img_ele = cel._element.xpath('.//pic:pic')[0]
embed = img_ele.xpath('.//a:blip/@r:embed')[0] #@表示获取器属性！加强学习！
# 下面是根据embed获取图片二进制数据
related_part = doc.part.related_parts[embed] #?????这里还没懂要看源码
image = related_part.image
image_bytes = image.blob #blob属性获取二进制
ext = image.ext #这里获取后缀,打印为png

tpl = DocxTemplate('123.docx') #模板导入成功
context = {
        'img':InlineImage(tpl, io.BytesIO(image_bytes))
    }
tpl.render(context)
tpl.save("FPGA反向生成的说明文档.docx")